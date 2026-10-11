# SPDX-License-Identifier: Apache-2.0
"""Offline synthetic JATS -> citation census -> verified writing facts canary.

Run with the promoted writing-documents Python and both source packages on
PYTHONPATH. This makes no scientific provider calls and does not test accuracy.
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from ocsl_review import census, literature
from ocsl_review.claim_inventory import build_claim_inventory
from ocsl_review.config import Paths
from ocsl_review.ingest import ingest_manuscript
from ocsl_review.util import atomic_json, sha256_file
from ocsl_writing import facts
from ocsl_writing.errors import ValidationError as WritingValidationError
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_literature import FakeScholar


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    workspace = args.output or Path(tempfile.mkdtemp(prefix="ocsl-census-writing-canary-"))
    workspace.mkdir(mode=0o700, parents=True, exist_ok=True)
    if any(workspace.iterdir()):
        raise RuntimeError("canary output must be a fresh, empty directory")
    manuscript_path = workspace / "synthetic-canary.xml"
    shutil.copyfile(root / "tests/fixtures/citation_census_canary.xml", manuscript_path)
    manuscript = ingest_manuscript(manuscript_path)
    assert manuscript["ingestion"]["complete"]
    atomic_json(workspace / ".ocsl-review/manuscript.json", manuscript, mode=0o600)
    graph = build_claim_inventory(manuscript)
    state, graph = literature.prepare_literature(workspace, manuscript, graph)
    state["confirmation"] = {"decision": "accepted", "at": "2026-10-11T00:00:00Z",
                             "scopeSha256": state["scopeSha256"],
                             "basis": "Operator-authorized synthetic offline integration test; no actual research or model judgment."}
    literature._save(workspace, state, graph, event="synthetic-canary-prepared")
    literature.fetch_literature(workspace, state, graph, FakeScholar())
    config = {"id": "synthetic-offline", "assignment": "Synthetic fixture only",
              "agent": "codex", "modelProvider": {"id": "synthetic", "model": "synthetic-fixture"},
              "providerId": "synthetic", "model": "synthetic-fixture", "bulkAvailable": True}
    command = SimpleNamespace(workspace=str(workspace), paths=Paths(root, None, None),
                              mode="full-text", worker_config=config["id"], reference_id=[],
                              limit=None, output="results/census")
    with patch.object(census, "_configuration", return_value=(0, config)):
        prepared = census.prepare(command, workspace, state, graph)
    command.run = prepared["runId"]
    records = []
    with census._database(workspace) as connection:
        tasks = census._task_rows(connection, command.run)
        for task in tasks:
            document = literature._verified_source(workspace, task["source"])
            records.append({"id": task["id"], "state": "completed", "providerId": "synthetic",
                            "model": "synthetic-fixture", "promptSha256": task["inputSha256"],
                            "usage": {"input": 100, "output": 20, "total": None, "quality": "partial"},
                            "output": {"supportVerdict": "scope-caveat", "supportConfidence": 50,
                              "rationale": "Synthetic observation used only to exercise validated mechanics.",
                              "readUnitIds": [block["id"] for block in document["blocks"]],
                              "evidence": [{"unitId": document["blocks"][0]["id"], "quote": document["blocks"][0]["text"]}],
                              "abstractSufficientReason": ""}})
        imported = census.import_results(command, workspace, connection, records)
    assert imported["accepted"] == len(tasks) and imported["rejected"] == 0
    exported = census.cmd_export(command)
    facts_path = Path(exported["facts"])
    verified = facts.verify_facts(facts_path)
    assert verified["ok"]
    assert verified["facts"]["citation_loci"]["value"] == len(tasks)
    assert verified["facts"]["calls_missing_total_tokens"]["value"] == len(tasks)
    template = workspace / "synthetic-report-template.md"
    template.write_text("# Synthetic integration canary\n\n"
                        "This is an offline fixture test, not scientific evidence or live-model validation.\n\n"
                        "Loci: {{fact:citation_loci}}; accepted observations: {{fact:model_assessed}}; "
                        "uncertain: {{fact:uncertain}}. Recorded input tokens: {{fact:input_tokens_recorded}}. "
                        "Calls with unknown totals: {{fact:calls_missing_total_tokens}}.\n", encoding="utf-8")
    report_path = workspace / "synthetic-report.md"
    rendered = facts.render_facts(facts_path, template, report_path)
    document = Document()
    for paragraph in report_path.read_text().split("\n\n"):
        document.add_paragraph(paragraph.strip())
    docx_path = workspace / "synthetic-report.docx"
    document.save(docx_path)
    assert "Loci: " + str(len(tasks)) in "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    tampered = workspace / "tampered-copy"
    shutil.copytree(facts_path.parent, tampered)
    with (tampered / "facts.sqlite3").open("ab") as handle:
        handle.write(b"tamper")
    try:
        facts.verify_facts(tampered / "facts.json")
    except WritingValidationError:
        tamper_rejected = True
    else:
        raise AssertionError("writing accepted a changed source hash")
    null_template = workspace / "unknown-total-template.md"
    null_template.write_text("Never invent a total: {{fact:total_tokens_recorded|d}}", encoding="utf-8")
    try:
        facts.render_facts(facts_path, null_template, workspace / "must-not-exist.md")
    except WritingValidationError:
        null_format_rejected = True
    else:
        raise AssertionError("writing formatted an unknown token count as a number")
    assert not (workspace / "must-not-exist.md").exists()
    report = {"ok": True, "synthetic": True, "liveModelCalls": 0,
              "workspace": str(workspace), "python": sys.executable,
              "runId": command.run, "loci": len(tasks), "bibliographyEntries": len(state["entries"]),
              "citationOccurrences": len(manuscript["citationClusters"]),
              "ownedAnchors": manuscript["structuredSource"]["ownedAnchors"],
              "facts": str(facts_path), "factsSha256": sha256_file(facts_path),
              "databaseSha256": sha256_file(facts_path.parent / "facts.sqlite3"),
              "rendered": rendered, "docx": str(docx_path), "docxSha256": sha256_file(docx_path),
              "tamperedSourceRejected": tamper_rejected, "unknownNumericTotalRejected": null_format_rejected,
              "meaning": "Mechanics and derivation only; no scientific validity or model-accuracy claim"}
    atomic_json(workspace / "canary-result.json", report, mode=0o600)
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
